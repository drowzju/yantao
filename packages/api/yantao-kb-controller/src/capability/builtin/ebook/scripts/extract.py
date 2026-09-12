# -*- coding: utf-8 -*-
"""
把一份资源文件抽成纯文本，输出一个 JSON 对象供 Node 侧读取。

用途：
  这是 yantao 工作台「读书项目」的取数脚本（ADR-0020 建立，ADR-0021 迁入能力
  目录）。它既被能力入口 `entry.py` 以函数方式调用（`extract_document`），也
  可以在命令行单独运行排查；命令行模式下把 resources/ 下的原始材料
  （pdf/epub/doc/docx/ppt/pptx/txt/md）抽成纯文本，stdout 除这一个 JSON 对象
  外不输出任何其它内容。

用法：
  python extract.py --format <fmt> --input <path>
    --format  文件格式：txt md pdf epub doc docx ppt pptx（由调用方按扩展名判定）
    --input   要抽取的文件的绝对路径
  示例：
    python extract.py --format pdf --input "D:\\kb\\resources\\三体.pdf"

依赖（按格式，缺哪个装哪个，均为运行时依赖）：
  txt/md   无
  pdf      pip install pypdf
  epub     pip install ebooklib（缺失时退回 zipfile+html 内置解析，阅读顺序可能不完美）
  docx     pip install python-docx
  pptx     pip install python-pptx
  doc      pip install pywin32 + 已安装 Word 桌面版（走 COM）
  ppt      pip install pywin32 + 已安装 PowerPoint 桌面版（走 COM）

退出码：
  0  成功
  3  找不到 --input 指定的文件
  4  缺少该格式所需的 Python 库
  5  文档没有可抽取的文字层（如扫描版 PDF）
  6  文档已加密或已损坏
  7  不支持的格式
  8  其它错误
  9  无法通过 Office COM 读取老格式（未装 Word/PowerPoint 或 COM 不可用）
  出错时 stdout 为空，stderr 输出 {"error": "中文说明", "kind": "..."}。

输出：
  {"text": "纯文本", "meta": {"format": "pdf", "chars": 12345}}
  text 统一 \n 换行；chars 是 text 的字符数。
"""

import argparse
import json
import os
import re
import sys
import zipfile
from html.parser import HTMLParser

EXIT_OK = 0
EXIT_INPUT = 3
EXIT_LIB = 4
EXIT_NO_TEXT = 5
EXIT_UNREADABLE = 6
EXIT_UNSUPPORTED = 7
EXIT_OTHER = 8
EXIT_OFFICE = 9

# kind → 命令行退出码；entry.py 只用 kind，main() 两个都用。
EXIT_BY_KIND = {
    "input-missing": EXIT_INPUT,
    "lib-missing": EXIT_LIB,
    "no-text": EXIT_NO_TEXT,
    "unreadable": EXIT_UNREADABLE,
    "unsupported": EXIT_UNSUPPORTED,
    "doc-unavailable": EXIT_OFFICE,
    "other": EXIT_OTHER,
}

FORMATS = ("txt", "md", "pdf", "epub", "doc", "docx", "ppt", "pptx")


class ExtractFailure(Exception):
    """一次可分类的抽取失败；kind 与 stderr JSON 的 kind 同词表。"""

    def __init__(self, kind, message):
        super().__init__(message)
        self.kind = kind
        self.message = message


def fail(message, kind, code):
    """把错误写成一行 JSON 到 stderr，然后以约定的退出码结束。"""
    sys.stderr.write(json.dumps({"error": message, "kind": kind}, ensure_ascii=False) + "\n")
    sys.stderr.flush()
    sys.exit(code)


def normalize(text):
    """统一换行并去掉首尾空白。"""
    return re.sub(r"\r\n?", "\n", text or "").strip()


class _TextFromHTML(HTMLParser):
    """粗糙的 HTML 剥标签器：跳过 script/style，其余文本节点收集成段落。"""

    def __init__(self):
        super().__init__()
        self.parts = []
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self._skip += 1

    def handle_endtag(self, tag):
        if tag in ("script", "style") and self._skip > 0:
            self._skip -= 1

    def handle_data(self, data):
        if self._skip == 0 and data.strip():
            self.parts.append(data.strip())


def strip_html(raw):
    """从 HTML 字节或字符串里剥出纯文本，段落之间空一行。"""
    if isinstance(raw, bytes):
        for encoding in ("utf-8", "gb18030"):
            try:
                raw = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raw = raw.decode("utf-8", errors="replace")
    parser = _TextFromHTML()
    parser.feed(raw)
    parser.close()
    return "\n\n".join(parser.parts)


def read_text_file(path):
    """txt/md 直读：依次尝试 utf-8（含 BOM）、utf-8、gb18030。"""
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            with open(path, "r", encoding=encoding) as handle:
                return handle.read()
        except UnicodeDecodeError:
            continue
    # 全都解不动就用替换模式兜底，不因几个坏字节丢掉整份文档。
    with open(path, "r", encoding="utf-8", errors="replace") as handle:
        return handle.read()


def extract_pdf(path):
    try:
        from pypdf import PdfReader
    except ImportError:
        raise ExtractFailure("lib-missing", "缺少 pypdf：请执行 pip install pypdf。") from None

    try:
        reader = PdfReader(path)
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ExtractFailure("unreadable", "这份 PDF 已加密，无法抽取文字。") from None
        pages = [(page.extract_text() or "") for page in reader.pages]
    except ExtractFailure:
        raise
    except Exception as error:
        raise ExtractFailure("unreadable", f"无法解析这份 PDF：{error}") from error

    text = "\n\n".join(pages)
    if not text.strip():
        raise ExtractFailure("no-text", "这份 PDF 没有可抽取的文字层（可能是扫描版）。")
    return text


def extract_epub(path):
    try:
        from ebooklib import epub, ITEM_DOCUMENT
    except ImportError:
        return extract_epub_zip(path)

    try:
        book = epub.read_epub(path)
        documents = list(book.get_items_of_type(ITEM_DOCUMENT))
    except Exception as error:
        raise ExtractFailure("unreadable", f"无法解析这份 EPUB：{error}") from error

    text = "\n\n".join(strip_html(item.get_content()) for item in documents)
    if not text.strip():
        raise ExtractFailure("no-text", "这份 EPUB 里没有可抽取的文字。")
    return text


def extract_epub_zip(path):
    """没有 ebooklib 时的兜底：把 zip 里所有 html/xhtml 按文件名排序剥标签。"""
    try:
        with zipfile.ZipFile(path) as archive:
            names = sorted(
                name for name in archive.namelist()
                if name.lower().endswith((".xhtml", ".html", ".htm"))
            )
            text = "\n\n".join(strip_html(archive.read(name)) for name in names)
    except Exception as error:
        raise ExtractFailure("unreadable", f"无法解析这份 EPUB：{error}") from error
    if not text.strip():
        raise ExtractFailure("no-text", "这份 EPUB 里没有可抽取的文字。")
    return text


def extract_docx(path):
    try:
        import docx
    except ImportError:
        raise ExtractFailure("lib-missing", "缺少 python-docx：请执行 pip install python-docx。") from None

    try:
        document = docx.Document(path)
    except Exception as error:
        raise ExtractFailure("unreadable", f"无法解析这份 DOCX：{error}") from error

    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def extract_pptx(path):
    try:
        from pptx import Presentation
    except ImportError:
        raise ExtractFailure("lib-missing", "缺少 python-pptx：请执行 pip install python-pptx。") from None

    try:
        presentation = Presentation(path)
    except Exception as error:
        raise ExtractFailure("unreadable", f"无法解析这份 PPTX：{error}") from error

    parts = []
    for index, slide in enumerate(presentation.slides, 1):
        parts.append(f"--- 第 {index} 页 ---")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in paragraph.runs)
                    if line.strip():
                        parts.append(line)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def office_com_text(application, document_path):
    """老 .doc/.ppt 共用的 COM 读法：打开 → 取文本 → 关闭，不弹窗口。"""
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    app = None
    try:
        app = win32com.client.DispatchEx(application)
        app.Visible = False
        if application == "Word.Application":
            document = app.Documents.Open(os.path.abspath(document_path), ReadOnly=True)
            try:
                # Word 的正文里段落是 \r，表格单元格结尾是 \x07，一并清掉。
                return normalize(document.Content.Text).replace("\x07", "")
            finally:
                document.Close(False)
        presentation = app.Presentations.Open(
            os.path.abspath(document_path), ReadOnly=True, Untitled=False, WithWindow=False
        )
        try:
            parts = []
            for index, slide in enumerate(presentation.Slides, 1):
                parts.append(f"--- 第 {index} 页 ---")
                for shape in slide.Shapes:
                    if shape.HasTextFrame:
                        text = normalize(shape.TextFrame.TextRange.Text)
                        if text:
                            parts.append(text)
            return "\n".join(parts)
        finally:
            presentation.Close()
    except ExtractFailure:
        raise
    except Exception as error:
        name = "Word" if application.startswith("Word") else "PowerPoint"
        raise ExtractFailure(
            "doc-unavailable",
            f"无法通过 {name} 读取这份文件：{error}。请确认已安装 {name} 桌面版。",
        ) from error
    finally:
        if app is not None:
            app.Quit()
        pythoncom.CoUninitialize()


def extract_doc(path):
    return office_com_text("Word.Application", path)


def extract_ppt(path):
    return office_com_text("PowerPoint.Application", path)


HANDLERS = {
    "txt": read_text_file,
    "md": read_text_file,
    "pdf": extract_pdf,
    "epub": extract_epub,
    "doc": extract_doc,
    "docx": extract_docx,
    "ppt": extract_ppt,
    "pptx": extract_pptx,
}


def extract_document(path, fmt):
    """
    抽取核心：把一份文件抽成统一换行的纯文本。供 entry.py 以函数方式调用；
    失败抛 ExtractFailure（kind/message 与命令行 stderr 的 JSON 同词表）。
    """
    if not os.path.isfile(path):
        raise ExtractFailure("input-missing", f"找不到要抽取的文件：{path}")

    handler = HANDLERS.get(fmt)
    if handler is None:
        raise ExtractFailure("unsupported", f"不支持的格式：{fmt}")

    try:
        return normalize(handler(path))
    except ExtractFailure:
        raise
    except Exception as error:
        raise ExtractFailure("other", f"抽取文档文本失败：{error}") from error


def main():
    parser = argparse.ArgumentParser(description="把一份资源文件抽成纯文本并输出 JSON")
    parser.add_argument("--format", required=True, choices=FORMATS, help="文件格式")
    parser.add_argument("--input", required=True, help="要抽取的文件路径")
    args = parser.parse_args()

    try:
        text = extract_document(args.input, args.format)
    except ExtractFailure as error:
        fail(error.message, error.kind, EXIT_BY_KIND.get(error.kind, EXIT_OTHER))
        return

    json.dump(
        {"text": text, "meta": {"format": args.format, "chars": len(text)}},
        sys.stdout,
        ensure_ascii=False,
    )
    sys.stdout.write("\n")
    sys.stdout.flush()
    sys.exit(EXIT_OK)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
    main()
