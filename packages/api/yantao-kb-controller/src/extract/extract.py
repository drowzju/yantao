# -*- coding: utf-8 -*-
"""
把一份资源文件抽成纯文本，输出一个 JSON 对象供 Node 侧读取。

用途：
  这是 yantao 工作台「读书项目」的取数脚本（ADR-0020）。它被 Node 以子进程
  方式调用，把 resources/ 下的原始材料（pdf/epub/doc/docx/ppt/pptx/txt/md）
  抽成纯文本；Node 把 stdout 里的 JSON 解析出来，缓存到知识库的
  `.yantao/extracts/` 下供 `kb_read_resource` 分页读取。stdout 除这一个
  JSON 对象外不输出任何其它内容。

用法：
  python extract.py --format <fmt> --input <path>
    --format  文件格式：txt md pdf epub doc docx ppt pptx（由 Node 按扩展名判定）
    --input   要抽取的文件的绝对路径
  示例：
    python extract.py --format pdf --input "D:\\kb\\resources\\三体.pdf"

依赖（按格式，缺哪个装哪个，均为运行时依赖）：
  txt/md   无
  pdf      pip install pypdf
  epub     pip install ebooklib（缺失时退回 zipfile+html 内置解析，阅读顺序可能不完美）
  docx     pip install python-docx
  pptx     pip install python-pptx
  doc      pip install pywin32 + 已安装 Word 桌面版（走 COM）
  ppt      pip install pywin32 + 已安装 PowerPoint 桌面版（走 COM）

退出码：
  0  成功
  3  找不到 --input 指定的文件
  4  缺少该格式所需的 Python 库
  5  文档没有可抽取的文字层（如扫描版 PDF）
  6  文档已加密或已损坏
  7  不支持的格式
  8  其它错误
  9  无法通过 Office COM 读取老格式（未装 Word/PowerPoint 或 COM 不可用）
  出错时 stdout 为空，stderr 输出 {"error": "中文说明", "kind": "..."}。

输出：
  {"text": "纯文本", "meta": {"format": "pdf", "chars": 12345}}
  text 统一 \n 换行；chars 是 text 的字符数。
"""

import argparse
import json
import os
import re
import sys
import zipfile
from html.parser import HTMLParser

EXIT_OK = 0
EXIT_INPUT = 3
EXIT_LIB = 4
EXIT_NO_TEXT = 5
EXIT_UNREADABLE = 6
EXIT_UNSUPPORTED = 7
EXIT_OTHER = 8
EXIT_OFFICE = 9

FORMATS = ("txt", "md", "pdf", "epub", "doc", "docx", "ppt", "pptx")


def fail(message, kind, code):
    """把错误写成一行 JSON 到 stderr，然后以约定的退出码结束。"""
    sys.stderr.write(json.dumps({"error": message, "kind": kind}, ensure_ascii=False) + "\n")
    sys.stderr.flush()
    sys.exit(code)


def normalize(text):
    """统一换行并去掉首尾空白。"""
    return re.sub(r"\r\n?", "\n", text or "").strip()


class _TextFromHTML(HTMLParser):
    """粗糙的 HTML 剥标签器：跳过 script/style，其余文本节点收集成段落。"""

    def __init__(self):
        super().__init__()
        self.parts = []
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self._skip += 1

    def handle_endtag(self, tag):
        if tag in ("script", "style") and self._skip > 0:
            self._skip -= 1

    def handle_data(self, data):
        if self._skip == 0 and data.strip():
            self.parts.append(data.strip())


def strip_html(raw):
    """从 HTML 字节或字符串里剥出纯文本，段落之间空一行。"""
    if isinstance(raw, bytes):
        for encoding in ("utf-8", "gb18030"):
            try:
                raw = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raw = raw.decode("utf-8", errors="replace")
    parser = _TextFromHTML()
    parser.feed(raw)
    parser.close()
    return "\n\n".join(parser.parts)


def read_text_file(path):
    """txt/md 直读：依次尝试 utf-8（含 BOM）、utf-8、gb18030。"""
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            with open(path, "r", encoding=encoding) as handle:
                return handle.read()
        except UnicodeDecodeError:
            continue
    # 全都解不动就用替换模式兜底，不因几个坏字节丢掉整份文档。
    with open(path, "r", encoding="utf-8", errors="replace") as handle:
        return handle.read()


def extract_pdf(path):
    try:
        from pypdf import PdfReader
    except ImportError:
        fail("缺少 pypdf：请执行 pip install pypdf。", "lib-missing", EXIT_LIB)
        return

    try:
        reader = PdfReader(path)
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                fail("这份 PDF 已加密，无法抽取文字。", "unreadable", EXIT_UNREADABLE)
                return
        pages = [(page.extract_text() or "") for page in reader.pages]
    except SystemExit:
        raise
    except Exception as error:
        fail(f"无法解析这份 PDF：{error}", "unreadable", EXIT_UNREADABLE)
        return

    text = "\n\n".join(pages)
    if not text.strip():
        fail("这份 PDF 没有可抽取的文字层（可能是扫描版）。", "no-text", EXIT_NO_TEXT)
        return
    return text


def extract_epub(path):
    try:
        from ebooklib import epub, ITEM_DOCUMENT
    except ImportError:
        return extract_epub_zip(path)

    try:
        book = epub.read_epub(path)
        documents = list(book.get_items_of_type(ITEM_DOCUMENT))
    except Exception as error:
        fail(f"无法解析这份 EPUB：{error}", "unreadable", EXIT_UNREADABLE)
        return

    text = "\n\n".join(strip_html(item.get_content()) for item in documents)
    if not text.strip():
        fail("这份 EPUB 里没有可抽取的文字。", "no-text", EXIT_NO_TEXT)
        return
    return text


def extract_epub_zip(path):
    """没有 ebooklib 时的兜底：把 zip 里所有 html/xhtml 按文件名排序剥标签。"""
    try:
        with zipfile.ZipFile(path) as archive:
            names = sorted(
                name for name in archive.namelist()
                if name.lower().endswith((".xhtml", ".html", ".htm"))
            )
            text = "\n\n".join(strip_html(archive.read(name)) for name in names)
    except Exception as error:
        fail(f"无法解析这份 EPUB：{error}", "unreadable", EXIT_UNREADABLE)
        return
    if not text.strip():
        fail("这份 EPUB 里没有可抽取的文字。", "no-text", EXIT_NO_TEXT)
        return
    return text


def extract_docx(path):
    try:
        import docx
    except ImportError:
        fail("缺少 python-docx：请执行 pip install python-docx。", "lib-missing", EXIT_LIB)
        return

    try:
        document = docx.Document(path)
    except Exception as error:
        fail(f"无法解析这份 DOCX：{error}", "unreadable", EXIT_UNREADABLE)
        return

    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def extract_pptx(path):
    try:
        from pptx import Presentation
    except ImportError:
        fail("缺少 python-pptx：请执行 pip install python-pptx。", "lib-missing", EXIT_LIB)
        return

    try:
        presentation = Presentation(path)
    except Exception as error:
        fail(f"无法解析这份 PPTX：{error}", "unreadable", EXIT_UNREADABLE)
        return

    parts = []
    for index, slide in enumerate(presentation.slides, 1):
        parts.append(f"--- 第 {index} 页 ---")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in paragraph.runs)
                    if line.strip():
                        parts.append(line)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def office_com_text(application, document_path):
    """老 .doc/.ppt 共用的 COM 读法：打开 → 取文本 → 关闭，不弹窗口。"""
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    app = None
    try:
        app = win32com.client.DispatchEx(application)
        app.Visible = False
        if application == "Word.Application":
            document = app.Documents.Open(os.path.abspath(document_path), ReadOnly=True)
            try:
                # Word 的正文里段落是 \r，表格单元格结尾是 \x07，一并清掉。
                return normalize(document.Content.Text).replace("\x07", "")
            finally:
                document.Close(False)
        presentation = app.Presentations.Open(
            os.path.abspath(document_path), ReadOnly=True, Untitled=False, WithWindow=False
        )
        try:
            parts = []
            for index, slide in enumerate(presentation.Slides, 1):
                parts.append(f"--- 第 {index} 页 ---")
                for shape in slide.Shapes:
                    if shape.HasTextFrame:
                        text = normalize(shape.TextFrame.TextRange.Text)
                        if text:
                            parts.append(text)
            return "\n".join(parts)
        finally:
            presentation.Close()
    except SystemExit:
        raise
    except Exception as error:
        name = "Word" if application.startswith("Word") else "PowerPoint"
        fail(
            f"无法通过 {name} 读取这份文件：{error}。请确认已安装 {name} 桌面版。",
            "doc-unavailable",
            EXIT_OFFICE,
        )
        return ""
    finally:
        if app is not None:
            app.Quit()
        pythoncom.CoUninitialize()


def extract_doc(path):
    return office_com_text("Word.Application", path)


def extract_ppt(path):
    return office_com_text("PowerPoint.Application", path)


HANDLERS = {
    "txt": read_text_file,
    "md": read_text_file,
    "pdf": extract_pdf,
    "epub": extract_epub,
    "doc": extract_doc,
    "docx": extract_docx,
    "ppt": extract_ppt,
    "pptx": extract_pptx,
}


def main():
    parser = argparse.ArgumentParser(description="把一份资源文件抽成纯文本并输出 JSON")
    parser.add_argument("--format", required=True, choices=FORMATS, help="文件格式")
    parser.add_argument("--input", required=True, help="要抽取的文件路径")
    args = parser.parse_args()

    if not os.path.isfile(args.input):
        fail(f"找不到要抽取的文件：{args.input}", "input-missing", EXIT_INPUT)
        return

    handler = HANDLERS.get(args.format)
    if handler is None:
        fail(f"不支持的格式：{args.format}", "unsupported", EXIT_UNSUPPORTED)
        return

    try:
        text = normalize(handler(args.input))
    except SystemExit:
        raise
    except Exception as error:
        fail(f"抽取文档文本失败：{error}", "other", EXIT_OTHER)
        return

    json.dump(
        {"text": text, "meta": {"format": args.format, "chars": len(text)}},
        sys.stdout,
        ensure_ascii=False,
    )
    sys.stdout.write("\n")
    sys.stdout.flush()
    sys.exit(EXIT_OK)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
    main()
